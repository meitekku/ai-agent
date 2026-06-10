"""Multi-format text extraction.

Dispatch function `extract_text(file_bytes, filename)` returns `(markdown, page_count)`.
"""

import asyncio
import csv
import io
import os
import re
from xml.etree.ElementTree import Element

from .ocr import ocr_pdf


def _escape_pipe(text: str) -> str:
    """Escape pipe characters in table cell text."""
    return text.replace("|", "\\|")


def _md_table(rows: list[list[str]]) -> str:
    """Build a Markdown table from rows (first row = header). Handles pipe escaping."""
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    lines = []
    # Header
    header = [_escape_pipe(c) for c in rows[0]] + [""] * (ncols - len(rows[0]))
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in range(ncols)) + " |")
    # Data
    for row in rows[1:]:
        padded = [_escape_pipe(c) for c in row] + [""] * (ncols - len(row))
        lines.append("| " + " | ".join(padded[:ncols]) + " |")
    return "\n".join(lines)


async def extract_text(file_bytes: bytes, filename: str) -> tuple[str, int]:
    """Extract text from various file formats.

    Returns (markdown_text, page_count).
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return await _extract_pdf(file_bytes, filename)
    elif ext in (".txt", ".md"):
        return _extract_text_plain(file_bytes)
    elif ext == ".csv":
        return await asyncio.to_thread(_extract_csv, file_bytes)
    elif ext == ".docx":
        return await asyncio.to_thread(_extract_docx, file_bytes)
    elif ext == ".xlsx":
        return await asyncio.to_thread(_extract_xlsx, file_bytes)
    elif ext == ".pptx":
        return await asyncio.to_thread(_extract_pptx, file_bytes)
    elif ext in (".html", ".htm"):
        return await asyncio.to_thread(_extract_html, file_bytes)
    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        return await _extract_image(file_bytes, filename, ext)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


async def _extract_pdf(file_bytes: bytes, filename: str) -> tuple[str, int]:
    pages = await ocr_pdf(file_bytes, filename)
    page_count = len(pages)
    markdown = "\n\n".join(
        (
            f"## Page {p['page']}\n\n{p['text']}"
            if not re.match(r"^#{1,6}\s", p["text"].strip())
            else p["text"]
        )
        for p in pages
    )
    return markdown, page_count


def _extract_text_plain(file_bytes: bytes) -> tuple[str, int]:
    text = file_bytes.decode("utf-8", errors="replace")
    return text, 1


def _decode_bytes(file_bytes: bytes) -> str:
    """Auto-detect encoding: try UTF-8 (with/without BOM), then cp932 (Shift_JIS superset)."""
    for enc in ("utf-8-sig", "utf-8", "cp932"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


# --- Structured CSV extraction (per-record documents) ---

_GROUP_KEY_RE = re.compile(r"番号|ID|No\.?$|コード|record.?id|key", re.IGNORECASE)
_NAME_RE = re.compile(r"名$|name|タイトル|件名|title|subject", re.IGNORECASE)
_DATE_RE = re.compile(r"日$|日時$|date|time", re.IGNORECASE)
_LONG_TEXT_THRESHOLD = 150


def _find_group_column(headers: list[str], data_rows: list[list[str]]) -> int | None:
    """Detect a column suitable for grouping rows (repeated ID-like values)."""
    candidates = [i for i, h in enumerate(headers) if _GROUP_KEY_RE.search(h)]
    for col in candidates:
        vals = [r[col].strip() for r in data_rows if col < len(r) and r[col].strip()]
        if len(vals) < 2:
            continue
        unique = set(vals)
        # Good grouping key: has repeated values but not all the same
        if len(unique) < len(vals) and len(unique) >= 2:
            return col
    return None


def _grouped_records(
    headers: list[str], data_rows: list[list[str]], group_col: int,
) -> str:
    """Build per-record documents grouped by a key column."""
    from collections import OrderedDict

    groups: OrderedDict[str, list[list[str]]] = OrderedDict()
    last_key = ""
    for row in data_rows:
        key = row[group_col].strip() if group_col < len(row) else ""
        if not key:
            key = last_key  # continuation row inherits previous key
        if not key:
            continue
        last_key = key
        groups.setdefault(key, []).append(row)

    # Find name/date columns for better titles
    name_col = next(
        (i for i, h in enumerate(headers) if _NAME_RE.search(h) and i != group_col),
        None,
    )
    date_col = next(
        (i for i, h in enumerate(headers) if _DATE_RE.search(h)), None,
    )

    sections: list[str] = []
    for key, rows in groups.items():
        # Classify columns: record-level (same across all rows) vs entry-level (varies)
        record_pairs: list[tuple[str, str]] = []
        entry_cols: list[int] = []

        for i, h in enumerate(headers):
            if i == group_col:
                continue
            vals = list(
                set(r[i].strip() for r in rows if i < len(r) and r[i].strip()),
            )
            if len(vals) == 0:
                continue
            elif len(vals) == 1:
                record_pairs.append((h, vals[0]))
            else:
                entry_cols.append(i)

        # Build section title
        name = ""
        if name_col is not None:
            name = next(
                (r[name_col].strip() for r in rows if name_col < len(r) and r[name_col].strip()),
                "",
            )
        title = f"## {headers[group_col]} {key}"
        if name:
            title += f" — {name}"

        lines: list[str] = [title]

        # Record-level fields (short → bullet list, long → subsection)
        for h, v in record_pairs:
            if len(v) <= _LONG_TEXT_THRESHOLD:
                lines.append(f"- {h}: {v}")
        for h, v in record_pairs:
            if len(v) > _LONG_TEXT_THRESHOLD:
                lines.append(f"\n### {h}\n{v}")

        # Entry-level data (one sub-entry per row)
        if entry_cols:
            has_multiple = len(rows) > 1
            if has_multiple:
                lines.append("\n### 活動・更新履歴")

            for row in rows:
                entry_short: list[str] = []
                entry_long: list[tuple[str, str]] = []
                date_val = ""

                for col in entry_cols:
                    if col >= len(row):
                        continue
                    val = row[col].strip()
                    if not val:
                        continue
                    h = headers[col]
                    if col == date_col:
                        date_val = val
                    if len(val) > _LONG_TEXT_THRESHOLD:
                        entry_long.append((h, val))
                    else:
                        entry_short.append(f"- {h}: {val}")

                if not entry_short and not entry_long:
                    continue

                entry_title = f"#### {date_val}" if date_val else "#### エントリ"
                lines.append(f"\n{entry_title}")
                lines.extend(entry_short)
                for h, v in entry_long:
                    lines.append(f"\n**{h}:**\n{v}")

        sections.append("\n".join(lines))

    return "\n\n---\n\n".join(sections)


def _flat_records(headers: list[str], data_rows: list[list[str]]) -> str:
    """Build per-row structured documents (no grouping)."""
    name_col = next(
        (i for i, h in enumerate(headers) if _NAME_RE.search(h)), 0,
    )

    sections: list[str] = []
    for i, row in enumerate(data_rows, 1):
        title = (
            row[name_col].strip()
            if name_col < len(row) and row[name_col].strip()
            else f"Record {i}"
        )

        lines: list[str] = [f"## {title}"]
        long_texts: list[tuple[str, str]] = []

        for j, h in enumerate(headers):
            if j >= len(row):
                break
            v = row[j].strip()
            if not v:
                continue
            if len(v) > _LONG_TEXT_THRESHOLD:
                long_texts.append((h, v))
            else:
                lines.append(f"- {h}: {v}")

        for h, v in long_texts:
            lines.append(f"\n### {h}\n{v}")

        if len(lines) > 1:  # has content beyond title
            sections.append("\n".join(lines))

    return "\n\n---\n\n".join(sections)


def _extract_csv(file_bytes: bytes) -> tuple[str, int]:
    text = _decode_bytes(file_bytes)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return "", 1

    headers = rows[0]
    data_rows = rows[1:]

    if not data_rows:
        return _md_table(rows), 1

    # Small simple tables: keep markdown table format
    if len(data_rows) <= 10 and len(headers) <= 8:
        return _md_table(rows), 1

    # Large/wide tables: per-record structured format
    group_col = _find_group_column(headers, data_rows)
    if group_col is not None:
        return _grouped_records(headers, data_rows, group_col), 1
    return _flat_records(headers, data_rows), 1


def _extract_docx(file_bytes: bytes) -> tuple[str, int]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    # Walk body elements in document order to preserve paragraph/table interleaving
    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

        if tag == "p":
            para = Paragraph(element, doc)
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                try:
                    level = int(para.style.name.replace("Heading", "").strip())
                    parts.append(f"{'#' * level} {text}")
                except ValueError:
                    parts.append(text)
            else:
                parts.append(text)

        elif tag == "tbl":
            table = Table(element, doc)
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                parts.append(_md_table(rows))

    return "\n\n".join(parts), 1


def _extract_xlsx(file_bytes: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):  # skip completely empty rows
                rows.append(cells)
        if not rows:
            continue

        headers = rows[0]
        data_rows = rows[1:]

        if not data_rows:
            parts.append(f"## {sheet.title}\n\n{_md_table(rows)}")
        elif len(data_rows) <= 10 and len(headers) <= 8:
            # Small simple tables: keep markdown table format
            parts.append(f"## {sheet.title}\n\n{_md_table(rows)}")
        else:
            # Large/wide tables: per-record structured format (same as CSV)
            group_col = _find_group_column(headers, data_rows)
            if group_col is not None:
                parts.append(_grouped_records(headers, data_rows, group_col))
            else:
                parts.append(_flat_records(headers, data_rows))

    wb.close()
    return "\n\n".join(parts), 1


def _extract_pptx(file_bytes: bytes) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    parts: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_texts.append(text)
            if shape.has_table:
                table = shape.table
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)
                if rows:
                    slide_texts.append(_md_table(rows))
        if slide_texts:
            parts.append(f"## Slide {i}\n\n" + "\n\n".join(slide_texts))

    return "\n\n".join(parts), len(prs.slides)


def _extract_html(file_bytes: bytes) -> tuple[str, int]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(file_bytes, "html.parser")

    # Remove script and style elements
    for tag in soup(["script", "style"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    # Clean up excessive whitespace
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)

    return text, 1


async def _extract_image(
    file_bytes: bytes, filename: str, ext: str
) -> tuple[str, int]:
    """Use Gemini Vision to extract text from images."""
    from . import config

    mime_map = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".webp": "image/webp",
    }
    mime_type = mime_map.get(ext, "image/png")

    from google import genai
    from google.genai import types

    client = genai.Client(api_key=config.GEMINI_API_KEY)
    resp = await client.aio.models.generate_content(
        model=config.GEMINI_OCR_MODEL,
        contents=[
            types.Part.from_bytes(data=file_bytes, mime_type=mime_type),
            "この画像の内容をMarkdown形式で正確に書き起こしてください。表はMarkdownテーブルに変換してください。テキストのみを出力してください。",
        ],
    )
    return resp.text or "", 1
